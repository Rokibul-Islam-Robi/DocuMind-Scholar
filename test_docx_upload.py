import docx
import urllib.request
import json
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

# Create a sample .docx document
doc = docx.Document()
doc.add_heading('Neuromorphic Spiking Transformers for Edge Robotics', level=1)
doc.add_paragraph('Abstract: Neuromorphic spiking transformers combine event-driven spike encoding with self-attention to achieve ultra-low power inference on neuromorphic silicon.')
doc.add_paragraph('Methodology: Membrane potential dynamics are modeled by the leaky integrate-and-fire equation: $$V[t] = \\beta V[t-1] + X[t] - S[t] V_{th}$$.')
doc.add_paragraph('Experiments: On DVS-Gesture benchmark, our model achieves 98.4% top-1 accuracy while drawing only 4.2 milliwatts of average power.')
doc.add_paragraph('Limitations: Spike timing jitter introduces latency variance in multi-hop neuromorphic mesh routing.')

docx_path = "sample_spiking_transformer.docx"
doc.save(docx_path)

# Upload the docx file
boundary = "----DocuMindBoundary5678"
with open(docx_path, "rb") as f:
    file_bytes = f.read()

body = []
body.append(f"--{boundary}".encode("utf-8"))
body.append(b'Content-Disposition: form-data; name="file"; filename="sample_spiking_transformer.docx"')
body.append(b'Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document')
body.append(b'')
body.append(file_bytes)
body.append(f"--{boundary}--".encode("utf-8"))
body.append(b'')
data = b"\r\n".join(body)

req = urllib.request.Request(
    "http://127.0.0.1:8000/api/papers/upload",
    data=data,
    headers={"Content-Type": f"multipart/form-data; boundary={boundary}"}
)

res = urllib.request.urlopen(req)
result = json.loads(res.read().decode("utf-8"))
print("DOCX Upload Result:", result["message"])
print("Parsed Title:", result["paper"]["title"], "| Format:", result["paper"]["file_type"])

# Clean up temporary test file
if os.path.exists(docx_path):
    os.remove(docx_path)
